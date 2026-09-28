from pathlib import Path
import json
import sys

import polars as pl
from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.config import MODEL_DIR, PROCESSED_DIR, ensure_directories
from src.data.event_definition import build_event_mapping_year
from src.data.survival_duration import build_survival_duration_year
from src.data.model_dataset import (
    START_YEAR,
    END_YEAR,
    CUTOFF,
    CORE_COLUMNS,
    build_model_dataset_year,
    validate_model_dataset_year,
    validate_model_frame,
)

YEARS = tuple(range(START_YEAR, END_YEAR + 1))
LABEL = f"{START_YEAR}_{END_YEAR}"


def count_rows(frame):
    return frame.select(pl.len()).collect().item()


def export_checked(frame, output):
    temporary = output.with_suffix(".tmp.parquet")
    frame.sink_parquet(temporary, compression="zstd")

    stats = validate_model_frame(pl.scan_parquet(temporary))
    if stats["rows"] == 0:
        raise ValueError(f"{output.name}: không có khoản vay nào.")

    temporary.replace(output)
    return stats


def export_excel(parquet_path):
    """Xuất toàn bộ bảng cấp khoản vay, có hiển thị tiến độ."""
    data = pl.read_parquet(parquet_path)
    output = parquet_path.with_suffix(".xlsx")
    temporary = output.with_name(output.stem + ".tmp.xlsx")

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Data_1")
    sheet.append(data.columns)

    # Trừ một dòng dành cho tiêu đề.
    limit = 1_048_575

    for index, row in enumerate(data.iter_rows()):
        if index > 0 and index % limit == 0:
            sheet = workbook.create_sheet(
                f"Data_{index // limit + 1}"
            )
            sheet.append(data.columns)

        sheet.append(row)

        if (index + 1) % 50000 == 0:
            print(
                f"Excel: đã ghi {index + 1:,}/{data.height:,} dòng...",
                flush=True,
            )

    print("Đang lưu file Excel...", flush=True)
    workbook.save(temporary)
    temporary.replace(output)
    print(f"Đã xuất Excel: {output}", flush=True)
    return output


def main():
    ensure_directories()
    report_dir = PROJECT_ROOT / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / f"model_dataset_validation_{LABEL}.json"

    report = {
        "cohort_years": list(YEARS),
        "performance_cutoff": CUTOFF,
        "unit": "one row per loan",
        "complete_case_variables": CORE_COLUMNS,
        "sample_flow_by_year": [],
        "status": "RUNNING",
    }

    def save_report():
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    save_report()

    try:
        orig_path = PROCESSED_DIR / "origination.parquet"

        required = [orig_path] + [
            MODEL_DIR / f"{kind}_clean_{year}.parquet"
            for year in YEARS
            for kind in ("orig", "perf")
        ]
        missing = [str(path) for path in required if not path.exists()]
        if missing:
            raise FileNotFoundError(
                "Thiếu file đầu vào:\n" + "\n".join(missing)
            )

        orig_all = pl.scan_parquet(orig_path)
        paths = []

        for year in YEARS:
            print(f"\nĐANG XỬ LÝ NĂM {year}", flush=True)

            orig_year = orig_all.filter(
                pl.col("vintage_year") == year
            )
            orig_count = count_rows(orig_year)
            if orig_count == 0:
                raise ValueError(
                    f"Origination đã chuẩn hóa không có năm {year}."
                )

            # Bắt buộc tạo lại theo logic event và thời gian mới.
            build_event_mapping_year(year, force=True)
            survival_path = build_survival_duration_year(
                year, force=True
            )
            model_path = build_model_dataset_year(year, force=True)
            validation = validate_model_dataset_year(year)

            survival = pl.scan_parquet(survival_path)
            model = pl.scan_parquet(model_path)

            no_followup = orig_year.select("loan_id").join(
                survival.select("loan_id"),
                on="loan_id",
                how="anti",
            )

            excluded = (
                survival.filter(
                    ~pl.col("survival_eligible").fill_null(False)
                )
                .group_by("survival_exclusion_reason")
                .len()
                .collect()
                .to_dicts()
            )

            event_counts = (
                model.group_by(
                    ["core_covariates_complete_flag", "event_type"]
                )
                .len()
                .collect()
                .to_dicts()
            )

            report["sample_flow_by_year"].append({
                "year": year,
                "origination_loans": orig_count,
                "without_usable_event_record": count_rows(no_followup),
                "time_exclusions": excluded,
                "analysis_loans": validation["rows"],
                "complete_case_loans": validation["complete_case"],
                "event_counts_by_completeness": event_counts,
            })
            save_report()
            paths.append(model_path)

        # Chỉ ghép khi toàn bộ 11 năm đều thành công.
        all_loans = pl.concat([
            pl.scan_parquet(path) for path in paths
        ]).sort("loan_id")

        analysis_path = MODEL_DIR / f"analysis_loans_{LABEL}.parquet"
        baseline_path = (
            MODEL_DIR / f"baseline_complete_cases_{LABEL}.parquet"
        )

        analysis_stats = export_checked(all_loans, analysis_path)

        baseline = all_loans.filter(
            pl.col("core_covariates_complete_flag")
        )
        baseline_stats = export_checked(baseline, baseline_path)

        report["analysis_validation"] = analysis_stats
        report["baseline_validation"] = baseline_stats
        report["status"] = "PARQUET_VALIDATED"
        save_report()

        # Excel gồm toàn bộ mẫu đủ 5 biến baseline.
        excel_path = export_excel(baseline_path)

        report["excel_file"] = str(excel_path)
        report["status"] = "PASS"
        save_report()

        print("\nHOÀN TẤT", flush=True)
        print(
            f"Đủ điều kiện thời gian: {analysis_stats['rows']:,} khoản"
        )
        print(
            f"Đủ biến baseline: {baseline_stats['rows']:,} khoản"
        )
        print(f"Excel: {excel_path}")
        print(f"Báo cáo: {report_path}")

    except Exception as error:
        report["status"] = "FAIL"
        report["error"] = str(error)
        save_report()
        raise


if __name__ == "__main__":
    main()